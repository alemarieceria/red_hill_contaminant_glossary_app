"""Validate immutable raw workbooks before canonical processing."""

from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path
from types import MappingProxyType

from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from .config import (
    FOOTNOTES_SHEET_NAME,
    GLOSSARY_SHEET_NAME,
    GLOSSARY_TABLE_NAME,
    GLOSSARY_WORKBOOK_FILENAME,
    METADATA_SHEET_NAME,
    METADATA_TABLE_NAME,
    REFERENCES_SHEET_NAME,
    REFERENCES_WORKBOOK_FILENAME,
    SUPPORTED_WORKBOOK_SCHEMA_VERSIONS,
    WORKBOOK_SCHEMA_VERSION,
    release_order_key,
    validate_release_id,
)
from .crosswalk import ReferenceResolutionMethod
from .footnotes import (
    FootnoteDefinition,
    FootnoteUsage,
    FootnoteValidationError,
    GlossaryFootnoteSource,
    validate_footnote_relationships,
)
from .intake import (
    IncomingContractError,
    IncomingPairInventory,
    IncomingWorkbookPair,
    IntakePublication,
    WorkbookInventory,
    inventory_incoming_pair,
    read_incoming_pair,
)
from .metadata import METADATA_HEADERS
from .paths import CONTAMINANT_REGISTRY_PATH, REFERENCE_CROSSWALK_PATH
from .registry_assets import (
    RegistryEntry,
    TrackedCrosswalkEntry,
    load_crosswalk,
    load_registry,
)
from .schemas import (
    FOOTNOTE_HEADER_MAP,
    GLOSSARY_HEADER_MAP,
    REFERENCE_HEADER_MAP,
)


class WorkbookContractError(ValueError):
    """Raised when an immutable raw workbook violates its schema contract."""


@dataclass(frozen=True)
class ValidatedWorkbookContract:
    """A completed intake whose raw workbook structure is schema compliant."""

    intake_publication: IntakePublication
    raw_pair: IncomingWorkbookPair
    raw_inventory: IncomingPairInventory
    data_release_id: str
    schema_version: str


class ValidationSeverity(StrEnum):
    """Severity assigned to one identity/relationship finding."""

    INFO = "info"
    WARNING = "warning"
    ERROR = "error"


class ValidationCategory(StrEnum):
    """Relationship area that produced a validation finding."""

    IDENTITY = "identity"
    REFERENCES = "references"
    FOOTNOTES = "footnotes"
    DUPLICATES = "duplicates"


@dataclass(frozen=True)
class ValidationFinding:
    """One deterministic problem or review note with source context."""

    code: str
    category: ValidationCategory
    severity: ValidationSeverity
    message: str
    workbook: str | None = None
    sheet: str | None = None
    source_row: int | None = None
    canonical_field: str | None = None
    id_contaminant: str | None = None
    source_value: str | None = None


@dataclass(frozen=True)
class ResolvedGlossaryIdentity:
    """One source glossary row resolved to its permanent contaminant ID."""

    source_row: int
    id_contaminant: str
    id_legacy_cg: int
    id_name: str
    id_casrn: object
    id_casrn_formula: str | None
    id_inchikey: object
    id_inchikey_formula: str | None
    footnote_value: object
    footnote_formula: str | None


@dataclass(frozen=True)
class ResolvedReferenceRelationship:
    """One source reference row joined through the reviewed crosswalk."""

    source_row: int
    refs_review_name: str
    id_contaminant: str
    resolution_method: ReferenceResolutionMethod


@dataclass(frozen=True)
class DuplicateIdentityCandidate:
    """One exact source token shared by multiple permanent identities."""

    canonical_field: str
    source_value: str
    id_contaminants: tuple[str, ...]
    source_rows: tuple[int, ...]


@dataclass(frozen=True)
class ValidatedIdentityRelationships:
    """Release-aware identity and relationship data from a valid contract."""

    workbook_contract: ValidatedWorkbookContract
    data_release_id: str
    schema_version: str
    registry_entries: tuple[RegistryEntry, ...]
    crosswalk_entries: tuple[TrackedCrosswalkEntry, ...]
    glossary_identities: tuple[ResolvedGlossaryIdentity, ...]
    reference_relationships: tuple[ResolvedReferenceRelationship, ...]
    footnote_definitions: tuple[FootnoteDefinition, ...]
    footnote_usages: tuple[FootnoteUsage, ...]
    duplicate_candidates: tuple[DuplicateIdentityCandidate, ...]
    findings: tuple[ValidationFinding, ...]


_SEVERITY_ORDER = MappingProxyType(
    {
        ValidationSeverity.ERROR: 0,
        ValidationSeverity.WARNING: 1,
        ValidationSeverity.INFO: 2,
    }
)
_CATEGORY_ORDER = MappingProxyType(
    {
        ValidationCategory.IDENTITY: 0,
        ValidationCategory.REFERENCES: 1,
        ValidationCategory.FOOTNOTES: 2,
        ValidationCategory.DUPLICATES: 3,
    }
)


def _finding_order_key(finding: ValidationFinding) -> tuple[object, ...]:
    return (
        _SEVERITY_ORDER[finding.severity],
        _CATEGORY_ORDER[finding.category],
        finding.code,
        finding.workbook or "",
        finding.sheet or "",
        finding.source_row if finding.source_row is not None else -1,
        finding.canonical_field or "",
        finding.id_contaminant or "",
        finding.source_value or "",
        finding.message,
    )


def _sorted_findings(
    findings: Iterable[ValidationFinding],
) -> tuple[ValidationFinding, ...]:
    return tuple(sorted(findings, key=_finding_order_key))


class IdentityRelationshipValidationError(ValueError):
    """Raised when identity/relationship validation has error findings."""

    def __init__(self, findings: Iterable[ValidationFinding]) -> None:
        self.findings = _sorted_findings(findings)
        message = "; ".join(
            f"{finding.code}: {finding.message}" for finding in self.findings
        )
        super().__init__(message)


@dataclass(frozen=True)
class _HeaderContract:
    workbook_role: str
    sheet_name: str
    table_name: str | None
    expected_headers: tuple[str, ...]


_SCHEMA_HEADER_CONTRACTS = MappingProxyType(
    {
        WORKBOOK_SCHEMA_VERSION: (
            _HeaderContract(
                "glossary",
                GLOSSARY_SHEET_NAME,
                GLOSSARY_TABLE_NAME,
                tuple(GLOSSARY_HEADER_MAP),
            ),
            _HeaderContract(
                "glossary",
                FOOTNOTES_SHEET_NAME,
                None,
                tuple(FOOTNOTE_HEADER_MAP),
            ),
            _HeaderContract(
                "glossary",
                METADATA_SHEET_NAME,
                METADATA_TABLE_NAME,
                METADATA_HEADERS,
            ),
            _HeaderContract(
                "references",
                REFERENCES_SHEET_NAME,
                None,
                tuple(REFERENCE_HEADER_MAP),
            ),
            _HeaderContract(
                "references",
                METADATA_SHEET_NAME,
                METADATA_TABLE_NAME,
                METADATA_HEADERS,
            ),
        )
    }
)


def _validated_publication(value: object) -> IntakePublication:
    if not isinstance(value, IntakePublication):
        raise WorkbookContractError(
            "workbook contract validation requires a completed IntakePublication"
        )

    try:
        release_id = validate_release_id(value.inventory.data_release_id)
    except (AttributeError, ValueError) as error:
        raise WorkbookContractError(
            "completed intake has an invalid data release ID"
        ) from error

    raw = value.raw_snapshot
    manifest = value.manifest
    if raw.inventory != value.inventory:
        raise WorkbookContractError(
            f"release {release_id} raw snapshot does not retain its accepted inventory"
        )
    if manifest.raw_snapshot != raw:
        raise WorkbookContractError(
            f"release {release_id} manifest does not retain its raw snapshot"
        )
    if (
        raw.data_release_id != release_id
        or manifest.manifest.data_release_id != release_id
        or raw.snapshot_dir.name != release_id
    ):
        raise WorkbookContractError(
            f"release {release_id} completed intake identities do not agree"
        )
    if (
        raw.glossary_path.parent != raw.snapshot_dir
        or raw.glossary_path.name != GLOSSARY_WORKBOOK_FILENAME
        or raw.references_path.parent != raw.snapshot_dir
        or raw.references_path.name != REFERENCES_WORKBOOK_FILENAME
    ):
        raise WorkbookContractError(
            f"release {release_id} raw workbook paths do not follow the stable contract"
        )
    return value


def _workbook_for_role(
    inventory: IncomingPairInventory,
    role: str,
) -> WorkbookInventory:
    if role == "glossary":
        return inventory.glossary_inventory
    if role == "references":
        return inventory.references_inventory
    raise WorkbookContractError(f"unsupported workbook contract role: {role!r}")


def _headers_for_contract(
    workbook: WorkbookInventory,
    contract: _HeaderContract,
) -> tuple[str, ...]:
    sheets = tuple(
        sheet for sheet in workbook.worksheets if sheet.name == contract.sheet_name
    )
    if len(sheets) != 1:
        raise WorkbookContractError(
            f"{contract.workbook_role} workbook expected exactly one sheet "
            f"{contract.sheet_name!r}; found {len(sheets)}"
        )
    sheet = sheets[0]
    if contract.table_name is None:
        headers = sheet.headers
    else:
        tables = tuple(
            table for table in sheet.tables if table.name == contract.table_name
        )
        if len(tables) != 1:
            raise WorkbookContractError(
                f"{contract.workbook_role} workbook sheet {contract.sheet_name!r} "
                f"expected exactly one table {contract.table_name!r}; "
                f"found {len(tables)}"
            )
        headers = tables[0].headers
    return tuple(header.value for header in headers)


def _validate_headers(
    *,
    release_id: str,
    inventory: IncomingPairInventory,
    contracts: tuple[_HeaderContract, ...],
) -> None:
    for contract in contracts:
        workbook = _workbook_for_role(inventory, contract.workbook_role)
        actual = _headers_for_contract(workbook, contract)
        expected = contract.expected_headers
        actual_counts = Counter(actual)
        expected_counts = Counter(expected)
        if actual_counts == expected_counts:
            continue

        missing = tuple(
            header
            for header in expected
            if actual_counts[header] < expected_counts[header]
        )
        unexpected = tuple(
            header
            for header in actual
            if actual_counts[header] > expected_counts[header]
        )
        location = (
            f"release {release_id} {contract.workbook_role} workbook "
            f"sheet {contract.sheet_name!r}"
        )
        if contract.table_name is not None:
            location += f" table {contract.table_name!r}"
        raise WorkbookContractError(
            f"{location} headers do not match schema: "
            f"missing {missing!r}; unexpected {unexpected!r}"
        )


def _require_matching_inventory(
    *,
    release_id: str,
    role: str,
    accepted: WorkbookInventory,
    observed: WorkbookInventory,
) -> None:
    if observed.size_bytes != accepted.size_bytes:
        raise WorkbookContractError(
            f"release {release_id} {role} raw snapshot byte size does not match "
            "the completed intake"
        )
    if observed.sha256 != accepted.sha256:
        raise WorkbookContractError(
            f"release {release_id} {role} raw snapshot SHA-256 does not match "
            "the completed intake"
        )
    if observed != accepted:
        raise WorkbookContractError(
            f"release {release_id} {role} raw snapshot structural inventory "
            "does not match the completed intake"
        )


def validate_workbook_contract(
    intake_publication: IntakePublication,
) -> ValidatedWorkbookContract:
    """Validate schema structure from a completed intake's raw snapshots."""

    completed = _validated_publication(intake_publication)
    release_id = completed.inventory.data_release_id
    raw_dir = completed.raw_snapshot.snapshot_dir
    try:
        raw_pair = read_incoming_pair(raw_dir)
        raw_inventory = inventory_incoming_pair(raw_pair)
    except IncomingContractError as error:
        raise WorkbookContractError(
            f"release {release_id} raw workbook contract failed: {error}"
        ) from error

    if raw_inventory.data_release_id != release_id:
        raise WorkbookContractError(
            f"release {release_id} raw workbook Metadata derives release "
            f"{raw_inventory.data_release_id!r}"
        )
    _require_matching_inventory(
        release_id=release_id,
        role="glossary",
        accepted=completed.inventory.glossary_inventory,
        observed=raw_inventory.glossary_inventory,
    )
    _require_matching_inventory(
        release_id=release_id,
        role="references",
        accepted=completed.inventory.references_inventory,
        observed=raw_inventory.references_inventory,
    )

    schema_version = raw_pair.compatibility.glossary_metadata.schema_version
    if schema_version not in SUPPORTED_WORKBOOK_SCHEMA_VERSIONS:
        raise WorkbookContractError(
            f"release {release_id} uses unsupported workbook schema "
            f"{schema_version!r}"
        )
    contracts = _SCHEMA_HEADER_CONTRACTS.get(schema_version)
    if contracts is None:
        raise WorkbookContractError(
            f"release {release_id} has no workbook header contract for supported "
            f"schema {schema_version!r}"
        )
    _validate_headers(
        release_id=release_id,
        inventory=raw_inventory,
        contracts=contracts,
    )

    return ValidatedWorkbookContract(
        intake_publication=completed,
        raw_pair=raw_pair,
        raw_inventory=raw_inventory,
        data_release_id=release_id,
        schema_version=schema_version,
    )


@dataclass(frozen=True)
class _GlossaryRelationshipSource:
    source_row: int
    values: Mapping[str, object]
    formulas: Mapping[str, str | None]


@dataclass(frozen=True)
class _ReferenceRelationshipSource:
    source_row: int
    value: object
    formula: str | None


def _source_header(header_map: Mapping[str, str], canonical_name: str) -> str:
    return next(
        header
        for header, mapped_name in header_map.items()
        if mapped_name == canonical_name
    )


def _sheet_cells(snapshot, sheet_name: str):
    sheet = next(sheet for sheet in snapshot.sheets if sheet.name == sheet_name)
    cells = {
        coordinate_to_tuple(cell.coordinate): cell for cell in sheet.cells
    }
    return sheet, cells


def _extract_relationship_glossary_rows(
    contract: ValidatedWorkbookContract,
) -> tuple[_GlossaryRelationshipSource, ...]:
    snapshot = contract.raw_pair.glossary_snapshot
    sheet, cells = _sheet_cells(snapshot, GLOSSARY_SHEET_NAME)
    table = next(table for table in sheet.tables if table.name == GLOSSARY_TABLE_NAME)
    min_column, header_row, max_column, max_row = range_boundaries(table.reference)
    header_columns = {
        cells[(header_row, column)].value: column
        for column in range(min_column, max_column + 1)
    }
    canonical_names = (
        "id_name",
        "id_legacy_cg",
        "id_casrn",
        "id_inchikey",
        "source_notes_footnote_ids",
    )
    columns = {
        canonical_name: header_columns[
            _source_header(GLOSSARY_HEADER_MAP, canonical_name)
        ]
        for canonical_name in canonical_names
    }
    rows = []
    for row in range(header_row + 1, max_row + 1):
        row_cells = {
            name: cells.get((row, column)) for name, column in columns.items()
        }
        rows.append(
            _GlossaryRelationshipSource(
                source_row=row,
                values=MappingProxyType(
                    {
                        name: cell.value if cell is not None else None
                        for name, cell in row_cells.items()
                    }
                ),
                formulas=MappingProxyType(
                    {
                        name: cell.formula if cell is not None else None
                        for name, cell in row_cells.items()
                    }
                ),
            )
        )
    return tuple(rows)


def _extract_relationship_reference_rows(
    contract: ValidatedWorkbookContract,
) -> tuple[_ReferenceRelationshipSource, ...]:
    snapshot = contract.raw_pair.references_snapshot
    sheet, cells = _sheet_cells(snapshot, REFERENCES_SHEET_NAME)
    header_columns = {
        cell.value: column
        for (row, column), cell in cells.items()
        if row == 1 and cell.value in REFERENCE_HEADER_MAP
    }
    label_column = header_columns[
        _source_header(REFERENCE_HEADER_MAP, "refs_review_name")
    ]
    mapped_columns = tuple(header_columns.values())
    rows = []
    for row in range(2, sheet.max_row + 1):
        if not any((row, column) in cells for column in mapped_columns):
            continue
        cell = cells.get((row, label_column))
        rows.append(
            _ReferenceRelationshipSource(
                source_row=row,
                value=cell.value if cell is not None else None,
                formula=cell.formula if cell is not None else None,
            )
        )
    return tuple(rows)


def _source_repr(value: object) -> str:
    return repr(value)


def _finding(
    code: str,
    category: ValidationCategory,
    severity: ValidationSeverity,
    message: str,
    *,
    workbook: str | None = None,
    sheet: str | None = None,
    source_row: int | None = None,
    canonical_field: str | None = None,
    id_contaminant: str | None = None,
    source_value: object = None,
    include_source_value: bool = False,
) -> ValidationFinding:
    return ValidationFinding(
        code=code,
        category=category,
        severity=severity,
        message=message,
        workbook=workbook,
        sheet=sheet,
        source_row=source_row,
        canonical_field=canonical_field,
        id_contaminant=id_contaminant,
        source_value=_source_repr(source_value) if include_source_value else None,
    )


def _applicable_registry_entries(
    entries: tuple[RegistryEntry, ...], release_id: str
) -> tuple[RegistryEntry, ...]:
    selected_key = release_order_key(release_id)
    return tuple(
        entry
        for entry in entries
        if release_order_key(entry.issued_release_id) <= selected_key
    )


def _entry_is_active(entry: RegistryEntry, release_id: str) -> bool:
    selected_key = release_order_key(release_id)
    return (
        release_order_key(entry.issued_release_id) <= selected_key
        and (
            entry.retired_release_id is None
            or release_order_key(entry.retired_release_id) > selected_key
        )
    )


def _resolve_glossary_identities(
    contract: ValidatedWorkbookContract,
    source_rows: tuple[_GlossaryRelationshipSource, ...],
    registry_entries: tuple[RegistryEntry, ...],
) -> tuple[tuple[ResolvedGlossaryIdentity, ...], tuple[ValidationFinding, ...]]:
    findings: list[ValidationFinding] = []
    workbook_name = contract.raw_pair.glossary_snapshot.path.name
    release_id = contract.data_release_id
    applicable = _applicable_registry_entries(registry_entries, release_id)
    active = tuple(
        entry for entry in applicable if _entry_is_active(entry, release_id)
    )
    applicable_by_legacy = {
        entry.id_legacy_cg: entry
        for entry in applicable
        if entry.id_legacy_cg is not None
    }
    active_ids = {entry.id_contaminant for entry in active}

    for entry in active:
        if entry.id_legacy_cg is None:
            findings.append(
                _finding(
                    "active_identity_without_legacy_id",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    f"active identity {entry.id_contaminant} cannot be represented "
                    f"by workbook schema {contract.schema_version}",
                    canonical_field="id_legacy_cg",
                    id_contaminant=entry.id_contaminant,
                )
            )

    parsed_rows: list[tuple[_GlossaryRelationshipSource, int, str]] = []
    seen_legacy: dict[int, int] = {}
    for source in source_rows:
        raw_legacy = source.values["id_legacy_cg"]
        raw_name = source.values["id_name"]
        row_valid = True
        if source.formulas["id_legacy_cg"] is not None:
            findings.append(
                _finding(
                    "formula_legacy_id",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    "legacy CG ID must be a literal positive integer",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_legacy_cg",
                    source_value=source.formulas["id_legacy_cg"],
                    include_source_value=True,
                )
            )
            row_valid = False
        elif type(raw_legacy) is not int or raw_legacy <= 0:
            findings.append(
                _finding(
                    "invalid_legacy_id",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    "legacy CG ID must be a literal positive integer",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_legacy_cg",
                    source_value=raw_legacy,
                    include_source_value=True,
                )
            )
            row_valid = False

        if source.formulas["id_name"] is not None:
            findings.append(
                _finding(
                    "formula_glossary_name",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    "glossary name must be literal nonblank text",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_name",
                    source_value=source.formulas["id_name"],
                    include_source_value=True,
                )
            )
            row_valid = False
        elif not isinstance(raw_name, str) or not raw_name.strip():
            findings.append(
                _finding(
                    "invalid_glossary_name",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    "glossary name must be literal nonblank text",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_name",
                    source_value=raw_name,
                    include_source_value=True,
                )
            )
            row_valid = False
        if not row_valid:
            continue

        assert type(raw_legacy) is int
        assert isinstance(raw_name, str)
        if raw_legacy in seen_legacy:
            findings.append(
                _finding(
                    "duplicate_legacy_id",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    f"legacy CG ID {raw_legacy} repeats source row "
                    f"{seen_legacy[raw_legacy]}",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_legacy_cg",
                    source_value=raw_legacy,
                    include_source_value=True,
                )
            )
        else:
            seen_legacy[raw_legacy] = source.source_row
        parsed_rows.append((source, raw_legacy, raw_name))

    resolved: list[ResolvedGlossaryIdentity] = []
    seen_contaminant: dict[str, int] = {}
    for source, legacy_id, name in parsed_rows:
        entry = applicable_by_legacy.get(legacy_id)
        if entry is None:
            findings.append(
                _finding(
                    "legacy_id_not_in_applicable_registry",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    f"legacy CG ID {legacy_id} has no registry identity applicable "
                    f"to release {release_id}",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_legacy_cg",
                    source_value=legacy_id,
                    include_source_value=True,
                )
            )
            continue
        if entry.id_contaminant not in active_ids:
            findings.append(
                _finding(
                    "retired_identity_present",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    f"retired identity {entry.id_contaminant} is present as a "
                    f"current glossary row for release {release_id}",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_legacy_cg",
                    id_contaminant=entry.id_contaminant,
                    source_value=legacy_id,
                    include_source_value=True,
                )
            )
            continue
        if entry.id_contaminant in seen_contaminant:
            findings.append(
                _finding(
                    "duplicate_resolved_identity",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    f"stable identity {entry.id_contaminant} repeats source row "
                    f"{seen_contaminant[entry.id_contaminant]}",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_contaminant",
                    id_contaminant=entry.id_contaminant,
                )
            )
        else:
            seen_contaminant[entry.id_contaminant] = source.source_row
        if name != entry.id_name:
            findings.append(
                _finding(
                    "registry_name_mismatch",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.WARNING,
                    f"glossary name {name!r} differs from registry name "
                    f"{entry.id_name!r}; stable identity is unchanged",
                    workbook=workbook_name,
                    sheet=GLOSSARY_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_name",
                    id_contaminant=entry.id_contaminant,
                    source_value=name,
                    include_source_value=True,
                )
            )
        resolved.append(
            ResolvedGlossaryIdentity(
                source_row=source.source_row,
                id_contaminant=entry.id_contaminant,
                id_legacy_cg=legacy_id,
                id_name=name,
                id_casrn=source.values["id_casrn"],
                id_casrn_formula=source.formulas["id_casrn"],
                id_inchikey=source.values["id_inchikey"],
                id_inchikey_formula=source.formulas["id_inchikey"],
                footnote_value=source.values["source_notes_footnote_ids"],
                footnote_formula=source.formulas["source_notes_footnote_ids"],
            )
        )

    present_ids = {identity.id_contaminant for identity in resolved}
    for entry in active:
        if entry.id_contaminant not in present_ids:
            findings.append(
                _finding(
                    "missing_active_identity",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    f"active identity {entry.id_contaminant} is absent from the "
                    f"glossary for release {release_id}",
                    canonical_field="id_contaminant",
                    id_contaminant=entry.id_contaminant,
                )
            )

    return (
        tuple(sorted(resolved, key=lambda value: value.id_contaminant)),
        tuple(findings),
    )


def _resolve_reference_relationships(
    contract: ValidatedWorkbookContract,
    source_rows: tuple[_ReferenceRelationshipSource, ...],
    crosswalk_entries: tuple[TrackedCrosswalkEntry, ...],
    identities: tuple[ResolvedGlossaryIdentity, ...],
    active_registry_ids: set[str],
) -> tuple[tuple[ResolvedReferenceRelationship, ...], tuple[ValidationFinding, ...]]:
    findings: list[ValidationFinding] = []
    workbook_name = contract.raw_pair.references_snapshot.path.name
    release_key = release_order_key(contract.data_release_id)
    eligible = tuple(
        entry
        for entry in crosswalk_entries
        if release_order_key(entry.reviewed_release_id) <= release_key
    )
    eligible_by_label = {entry.refs_review_name: entry for entry in eligible}
    identities_by_id = {entry.id_contaminant: entry for entry in identities}
    present_ids = set(identities_by_id)
    resolved = []
    mismatched_labels: set[str] = set()

    for source in source_rows:
        if source.formula is not None:
            findings.append(
                _finding(
                    "formula_reference_label",
                    ValidationCategory.REFERENCES,
                    ValidationSeverity.ERROR,
                    "reference label must be literal nonblank text",
                    workbook=workbook_name,
                    sheet=REFERENCES_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="refs_review_name",
                    source_value=source.formula,
                    include_source_value=True,
                )
            )
            continue
        if not isinstance(source.value, str) or not source.value.strip():
            findings.append(
                _finding(
                    "invalid_reference_label",
                    ValidationCategory.REFERENCES,
                    ValidationSeverity.ERROR,
                    "reference label must be literal nonblank text",
                    workbook=workbook_name,
                    sheet=REFERENCES_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="refs_review_name",
                    source_value=source.value,
                    include_source_value=True,
                )
            )
            continue
        label = source.value
        entry = eligible_by_label.get(label)
        if entry is None:
            findings.append(
                _finding(
                    "unresolved_reference_label",
                    ValidationCategory.REFERENCES,
                    ValidationSeverity.ERROR,
                    f"reference label {label!r} has no reviewed crosswalk entry "
                    f"available for release {contract.data_release_id}",
                    workbook=workbook_name,
                    sheet=REFERENCES_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="refs_review_name",
                    source_value=label,
                    include_source_value=True,
                )
            )
            continue
        if entry.id_contaminant not in active_registry_ids:
            findings.append(
                _finding(
                    "reference_target_not_active",
                    ValidationCategory.REFERENCES,
                    ValidationSeverity.ERROR,
                    f"crosswalk target {entry.id_contaminant} is not active at "
                    f"release {contract.data_release_id}",
                    workbook=workbook_name,
                    sheet=REFERENCES_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_contaminant",
                    id_contaminant=entry.id_contaminant,
                    source_value=label,
                    include_source_value=True,
                )
            )
            continue
        if entry.id_contaminant not in present_ids:
            findings.append(
                _finding(
                    "reference_target_absent_from_glossary",
                    ValidationCategory.REFERENCES,
                    ValidationSeverity.ERROR,
                    f"crosswalk target {entry.id_contaminant} is absent from the "
                    "selected glossary",
                    workbook=workbook_name,
                    sheet=REFERENCES_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="id_contaminant",
                    id_contaminant=entry.id_contaminant,
                    source_value=label,
                    include_source_value=True,
                )
            )
            continue
        resolved.append(
            ResolvedReferenceRelationship(
                source_row=source.source_row,
                refs_review_name=label,
                id_contaminant=entry.id_contaminant,
                resolution_method=entry.resolution_method,
            )
        )
        if (
            label != identities_by_id[entry.id_contaminant].id_name
            and label not in mismatched_labels
        ):
            mismatched_labels.add(label)
            findings.append(
                _finding(
                    "reference_label_differs_from_glossary_name",
                    ValidationCategory.REFERENCES,
                    ValidationSeverity.INFO,
                    f"reviewed label {label!r} differs from target glossary name "
                    f"{identities_by_id[entry.id_contaminant].id_name!r}",
                    workbook=workbook_name,
                    sheet=REFERENCES_SHEET_NAME,
                    source_row=source.source_row,
                    canonical_field="refs_review_name",
                    id_contaminant=entry.id_contaminant,
                    source_value=label,
                    include_source_value=True,
                )
            )

    return tuple(resolved), tuple(findings)


def _validate_resolved_footnotes(
    contract: ValidatedWorkbookContract,
    identities: tuple[ResolvedGlossaryIdentity, ...],
) -> tuple[
    tuple[FootnoteDefinition, ...],
    tuple[FootnoteUsage, ...],
    tuple[ValidationFinding, ...],
]:
    unique_sources: dict[str, GlossaryFootnoteSource] = {}
    for identity in identities:
        unique_sources.setdefault(
            identity.id_contaminant,
            GlossaryFootnoteSource(
                id_contaminant=identity.id_contaminant,
                source_row=identity.source_row,
                value=identity.footnote_value,
                formula=identity.footnote_formula,
            ),
        )
    try:
        definitions, usages = validate_footnote_relationships(
            contract.raw_pair.glossary_snapshot,
            unique_sources.values(),
        )
    except FootnoteValidationError as error:
        workbook_name = contract.raw_pair.glossary_snapshot.path.name
        findings = tuple(
            _finding(
                "invalid_footnote_relationship",
                ValidationCategory.FOOTNOTES,
                ValidationSeverity.ERROR,
                issue.message,
                workbook=workbook_name,
                sheet=issue.sheet_name,
                source_row=issue.source_row,
                canonical_field=(
                    "source_notes_footnote_ids"
                    if issue.sheet_name == GLOSSARY_SHEET_NAME
                    else "source_notes_footnote_id"
                ),
                source_value=issue.source_value,
                include_source_value=issue.source_value is not None,
            )
            for issue in error.issues
        )
        return (), (), findings
    return definitions, usages, ()


def _duplicate_candidates(
    identities: tuple[ResolvedGlossaryIdentity, ...],
) -> tuple[tuple[DuplicateIdentityCandidate, ...], tuple[ValidationFinding, ...]]:
    grouped: dict[tuple[str, str], dict[str, set[int]]] = defaultdict(
        lambda: defaultdict(set)
    )
    for identity in identities:
        grouped[("id_name", identity.id_name)][identity.id_contaminant].add(
            identity.source_row
        )
        for canonical_field, raw_value in (
            ("id_casrn", identity.id_casrn),
            ("id_inchikey", identity.id_inchikey),
        ):
            formula = (
                identity.id_casrn_formula
                if canonical_field == "id_casrn"
                else identity.id_inchikey_formula
            )
            if formula is not None:
                continue
            if not isinstance(raw_value, str):
                continue
            for token in set(raw_value.split(" | ")):
                if token in {"NA", "N/A"}:
                    continue
                grouped[(canonical_field, token)][identity.id_contaminant].add(
                    identity.source_row
                )

    candidates = []
    findings = []
    code_by_field = {
        "id_name": "duplicate_name_candidate",
        "id_casrn": "duplicate_casrn_candidate",
        "id_inchikey": "duplicate_inchikey_candidate",
    }
    for (canonical_field, source_value), rows_by_id in grouped.items():
        if len(rows_by_id) < 2:
            continue
        ids = tuple(sorted(rows_by_id))
        rows = tuple(sorted({row for values in rows_by_id.values() for row in values}))
        candidate = DuplicateIdentityCandidate(
            canonical_field=canonical_field,
            source_value=source_value,
            id_contaminants=ids,
            source_rows=rows,
        )
        candidates.append(candidate)
        findings.append(
            _finding(
                code_by_field[canonical_field],
                ValidationCategory.DUPLICATES,
                ValidationSeverity.WARNING,
                f"exact {canonical_field} value {source_value!r} is shared by "
                f"distinct identities {', '.join(ids)}; records were not merged",
                canonical_field=canonical_field,
                source_value=source_value,
                include_source_value=True,
            )
        )
    candidates.sort(
        key=lambda value: (
            value.canonical_field,
            value.source_value,
            value.id_contaminants,
            value.source_rows,
        )
    )
    return tuple(candidates), tuple(findings)


def validate_identity_relationships(
    validated_contract: ValidatedWorkbookContract,
    registry_path: Path = CONTAMINANT_REGISTRY_PATH,
    crosswalk_path: Path = REFERENCE_CROSSWALK_PATH,
) -> ValidatedIdentityRelationships:
    """Validate release-aware IDs and joins from immutable raw snapshots."""

    if not isinstance(validated_contract, ValidatedWorkbookContract):
        raise IdentityRelationshipValidationError(
            (
                _finding(
                    "invalid_workbook_contract",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    "identity validation requires a ValidatedWorkbookContract",
                ),
            )
        )
    if validated_contract.schema_version != WORKBOOK_SCHEMA_VERSION:
        raise IdentityRelationshipValidationError(
            (
                _finding(
                    "unsupported_identity_schema",
                    ValidationCategory.IDENTITY,
                    ValidationSeverity.ERROR,
                    f"identity validation has no rules for schema "
                    f"{validated_contract.schema_version!r}",
                ),
            )
        )

    findings: list[ValidationFinding] = []
    registry_entries: tuple[RegistryEntry, ...] = ()
    crosswalk_entries: tuple[TrackedCrosswalkEntry, ...] = ()
    registry_loaded = False
    crosswalk_loaded = False
    try:
        registry_entries = load_registry(Path(registry_path))
        registry_loaded = True
    except (OSError, ValueError) as error:
        findings.append(
            _finding(
                "invalid_registry_asset",
                ValidationCategory.IDENTITY,
                ValidationSeverity.ERROR,
                f"could not load contaminant registry: {error}",
                workbook=Path(registry_path).name,
            )
        )
    if registry_loaded:
        try:
            crosswalk_entries = load_crosswalk(
                Path(crosswalk_path), registry_entries
            )
            crosswalk_loaded = True
        except (OSError, ValueError) as error:
            findings.append(
                _finding(
                    "invalid_crosswalk_asset",
                    ValidationCategory.REFERENCES,
                    ValidationSeverity.ERROR,
                    f"could not load reference crosswalk: {error}",
                    workbook=Path(crosswalk_path).name,
                )
            )

    try:
        glossary_sources = _extract_relationship_glossary_rows(validated_contract)
        reference_sources = _extract_relationship_reference_rows(validated_contract)
    except (KeyError, StopIteration, ValueError) as error:
        findings.append(
            _finding(
                "invalid_relationship_source_structure",
                ValidationCategory.IDENTITY,
                ValidationSeverity.ERROR,
                f"validated workbook structure cannot be read: {error}",
            )
        )
        glossary_sources = ()
        reference_sources = ()

    identities: tuple[ResolvedGlossaryIdentity, ...] = ()
    if registry_loaded:
        identities, identity_findings = _resolve_glossary_identities(
            validated_contract,
            glossary_sources,
            registry_entries,
        )
        findings.extend(identity_findings)

    release_id = validated_contract.data_release_id
    applicable_registry = _applicable_registry_entries(
        registry_entries, release_id
    )
    active_registry_ids = {
        entry.id_contaminant
        for entry in applicable_registry
        if _entry_is_active(entry, release_id)
    }
    eligible_crosswalk = tuple(
        entry
        for entry in crosswalk_entries
        if release_order_key(entry.reviewed_release_id)
        <= release_order_key(release_id)
    )
    relationships: tuple[ResolvedReferenceRelationship, ...] = ()
    if crosswalk_loaded:
        relationships, reference_findings = _resolve_reference_relationships(
            validated_contract,
            reference_sources,
            crosswalk_entries,
            identities,
            active_registry_ids,
        )
        findings.extend(reference_findings)

    definitions, usages, footnote_findings = _validate_resolved_footnotes(
        validated_contract, identities
    )
    findings.extend(footnote_findings)
    candidates, duplicate_findings = _duplicate_candidates(identities)
    findings.extend(duplicate_findings)

    sorted_findings = _sorted_findings(findings)
    if any(
        finding.severity is ValidationSeverity.ERROR
        for finding in sorted_findings
    ):
        raise IdentityRelationshipValidationError(sorted_findings)

    return ValidatedIdentityRelationships(
        workbook_contract=validated_contract,
        data_release_id=release_id,
        schema_version=validated_contract.schema_version,
        registry_entries=applicable_registry,
        crosswalk_entries=eligible_crosswalk,
        glossary_identities=identities,
        reference_relationships=relationships,
        footnote_definitions=definitions,
        footnote_usages=usages,
        duplicate_candidates=candidates,
        findings=sorted_findings,
    )
