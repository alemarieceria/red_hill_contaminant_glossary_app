"""Extract and validate authoritative workbook Metadata."""

from dataclasses import dataclass
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from .config import (
    GLOSSARY_WORKBOOK_TYPE,
    METADATA_SHEET_NAME,
    METADATA_TABLE_NAME,
    REFERENCES_WORKBOOK_TYPE,
    SUPPORTED_WORKBOOK_SCHEMA_VERSIONS,
    release_order_key,
    validate_release_id,
)
from .io_excel import WorkbookSnapshot


METADATA_HEADERS = ("key", "value")
METADATA_KEYS = frozenset(
    {"workbook_type", "schema_version", "workbook_revision"}
)


@dataclass(frozen=True)
class WorkbookMetadata:
    """The three contract values declared by one workbook."""

    workbook_type: str
    schema_version: str
    workbook_revision: str


@dataclass(frozen=True)
class WorkbookCompatibility:
    """A validated workbook pair and its combined release identifier."""

    glossary_metadata: WorkbookMetadata
    references_metadata: WorkbookMetadata
    data_release_id: str


def _literal_text(value: object, field_name: str) -> str:
    """Return exact nonblank text or reject normalization-dependent values."""

    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Metadata {field_name} must be nonblank text")
    if value != value.strip():
        raise ValueError(
            f"Metadata {field_name} must not have surrounding whitespace"
        )
    return value


def _extract_workbook_metadata(snapshot: WorkbookSnapshot) -> WorkbookMetadata:
    """Implement Metadata extraction before workbook context is attached."""

    metadata_sheets = tuple(
        sheet for sheet in snapshot.sheets if sheet.name == METADATA_SHEET_NAME
    )
    if len(metadata_sheets) != 1:
        raise ValueError(
            f"{snapshot.path}: expected exactly one {METADATA_SHEET_NAME!r} "
            f"sheet; found {len(metadata_sheets)}"
        )
    sheet = metadata_sheets[0]

    if len(sheet.tables) != 1 or sheet.tables[0].name != METADATA_TABLE_NAME:
        table_names = ", ".join(repr(table.name) for table in sheet.tables)
        raise ValueError(
            f"{snapshot.path}: expected exactly one {METADATA_TABLE_NAME!r} "
            f"table on {METADATA_SHEET_NAME!r}; found [{table_names}]"
        )
    table = sheet.tables[0]

    try:
        min_column, header_row, max_column, max_row = range_boundaries(
            table.reference
        )
    except ValueError as error:
        raise ValueError(
            f"{snapshot.path}: malformed MetadataTable range: "
            f"{table.reference!r}"
        ) from error

    if max_column - min_column + 1 != 2:
        raise ValueError("MetadataTable must contain exactly two columns")
    if max_row - header_row + 1 != 4:
        raise ValueError(
            "MetadataTable must contain one header row and three data rows"
        )

    cells = {
        coordinate_to_tuple(cell.coordinate): cell for cell in sheet.cells
    }
    table_cells = tuple(
        cells.get((row, column))
        for row in range(header_row, max_row + 1)
        for column in range(min_column, max_column + 1)
    )
    if any(cell is not None and cell.formula is not None for cell in table_cells):
        raise ValueError("MetadataTable values must be literal text, not formulas")

    headers = tuple(
        cells.get((header_row, column)).value
        if cells.get((header_row, column)) is not None
        else None
        for column in range(min_column, max_column + 1)
    )
    if headers != METADATA_HEADERS:
        raise ValueError(
            f"MetadataTable headers must be exactly {METADATA_HEADERS!r}; "
            f"found {headers!r}"
        )

    values_by_key: dict[str, str] = {}
    duplicate_keys: set[str] = set()
    for row in range(header_row + 1, max_row + 1):
        key_cell = cells.get((row, min_column))
        value_cell = cells.get((row, min_column + 1))
        key = _literal_text(
            key_cell.value if key_cell is not None else None,
            "key",
        )
        value = _literal_text(
            value_cell.value if value_cell is not None else None,
            f"value for {key!r}",
        )
        if key in values_by_key:
            duplicate_keys.add(key)
        else:
            values_by_key[key] = value

    if duplicate_keys:
        duplicates = ", ".join(repr(key) for key in sorted(duplicate_keys))
        raise ValueError(f"duplicate Metadata keys: {duplicates}")

    actual_keys = set(values_by_key)
    missing_keys = sorted(METADATA_KEYS - actual_keys)
    unknown_keys = sorted(actual_keys - METADATA_KEYS)
    if missing_keys or unknown_keys:
        details: list[str] = []
        if missing_keys:
            details.append(
                "missing " + ", ".join(repr(key) for key in missing_keys)
            )
        if unknown_keys:
            details.append(
                "unknown " + ", ".join(repr(key) for key in unknown_keys)
            )
        raise ValueError("invalid Metadata keys: " + "; ".join(details))

    revision = values_by_key["workbook_revision"]
    validate_release_id(revision)
    return WorkbookMetadata(
        workbook_type=values_by_key["workbook_type"],
        schema_version=values_by_key["schema_version"],
        workbook_revision=revision,
    )


def extract_workbook_metadata(snapshot: WorkbookSnapshot) -> WorkbookMetadata:
    """Extract validated Metadata from an in-memory workbook snapshot."""

    try:
        return _extract_workbook_metadata(snapshot)
    except ValueError as error:
        raise ValueError(
            f"{snapshot.path} {METADATA_TABLE_NAME}: {error}"
        ) from error


def validate_workbook_compatibility(
    glossary_metadata: WorkbookMetadata,
    references_metadata: WorkbookMetadata,
) -> WorkbookCompatibility:
    """Validate a glossary/references pair and derive its release ID."""

    if not isinstance(glossary_metadata, WorkbookMetadata):
        raise ValueError("glossary Metadata must be a WorkbookMetadata record")
    if not isinstance(references_metadata, WorkbookMetadata):
        raise ValueError("references Metadata must be a WorkbookMetadata record")

    if glossary_metadata.workbook_type != GLOSSARY_WORKBOOK_TYPE:
        raise ValueError(
            "glossary workbook must declare workbook_type "
            f"{GLOSSARY_WORKBOOK_TYPE!r}"
        )
    if references_metadata.workbook_type != REFERENCES_WORKBOOK_TYPE:
        raise ValueError(
            "references workbook must declare workbook_type "
            f"{REFERENCES_WORKBOOK_TYPE!r}"
        )

    glossary_schema = _literal_text(
        glossary_metadata.schema_version, "schema_version"
    )
    references_schema = _literal_text(
        references_metadata.schema_version, "schema_version"
    )
    if glossary_schema != references_schema:
        raise ValueError(
            "workbook schema versions do not match: "
            f"{glossary_schema!r} != {references_schema!r}"
        )
    if glossary_schema not in SUPPORTED_WORKBOOK_SCHEMA_VERSIONS:
        raise ValueError(f"unsupported workbook schema version: {glossary_schema}")

    glossary_revision = glossary_metadata.workbook_revision
    references_revision = references_metadata.workbook_revision
    release_id = max(
        (glossary_revision, references_revision),
        key=release_order_key,
    )
    return WorkbookCompatibility(
        glossary_metadata=glossary_metadata,
        references_metadata=references_metadata,
        data_release_id=release_id,
    )
